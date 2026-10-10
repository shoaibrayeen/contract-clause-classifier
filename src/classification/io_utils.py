"""Read contract text from .txt / .md / .pdf / .docx files or raw bytes."""

from __future__ import annotations

import io
from pathlib import Path

TEXT_SUFFIXES = {".txt", ".md", ".text", ""}


def read_pdf_bytes(data: bytes) -> str:
    import pymupdf  # PyMuPDF

    with pymupdf.open(stream=data, filetype="pdf") as doc:
        text = "\n\n".join(page.get_text("text") for page in doc)
    if not text.strip():
        raise ValueError("PDF has no text layer (scanned PDFs need OCR first)")
    return text


def read_docx_bytes(data: bytes) -> str:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n\n".join(p for p in parts if p.strip())


def read_bytes(data: bytes, filename: str = "") -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf" or data[:5] == b"%PDF-":
        return read_pdf_bytes(data)
    if suffix == ".docx" or data[:2] == b"PK" and suffix not in TEXT_SUFFIXES:
        return read_docx_bytes(data)
    return data.decode("utf-8", errors="replace")


def read_file(path: str | Path) -> str:
    path = Path(path)
    return read_bytes(path.read_bytes(), path.name)
